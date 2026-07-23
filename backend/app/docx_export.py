"""Εξαγωγή υποδείγματος σε .docx με την επιστολόχαρτη μορφή του γραφείου.

Επιστολόχαρτο Skotanis & Associates: χρυσή γραμμή-σήμα πάνω από την επωνυμία,
επωνυμία σε navy serif, υπότιτλος «Elite Legal Advocacy» σε χρυσό, χρυσός
κανόνας κάτω από την κεφαλίδα· υποσέλιδο με τα στοιχεία επικοινωνίας των
δύο γραφείων. Το σώμα αποδίδεται με τους κανόνες στοιχειοθεσίας δικογράφου:
κεντραρισμένες κεφαλίδες, πλήρης στοίχιση παραγράφων, δεξιά στοίχιση του
μπλοκ υπογραφής.
"""

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.exemplars import Exemplar

NAVY = RGBColor(0x0A, 0x16, 0x28)
GOLD = RGBColor(0xC6, 0xA7, 0x5E)
GOLD_DARK = RGBColor(0xA8, 0x89, 0x3A)
SILVER = RGBColor(0x6B, 0x77, 0x86)

BODY_FONT = "Georgia"
DISPLAY_FONT = "Cormorant Garamond"

_BRACKETS = re.compile(r"\[[^\]]*\]")


def _set_font(run, *, name: str, size: float, color: RGBColor, bold: bool = False,
              italic: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    # Ρητή δήλωση hAnsi ώστε τα ελληνικά να αποδίδονται με την ίδια γραμματοσειρά.
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None:
        rfonts.set(qn("w:hAnsi"), name)


def _add_border(paragraph, edge: str, *, color: str, size: int = 6,
                space: int = 4) -> None:
    """Προσθέτει γραμμή-περίγραμμα (w:pBdr) σε πλευρά παραγράφου."""
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = ppr.makeelement(qn("w:pBdr"), {})
        ppr.append(pbdr)
    element = pbdr.makeelement(qn(f"w:{edge}"), {})
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), str(space))
    element.set(qn("w:color"), color)
    pbdr.append(element)


def _is_display_line(line: str) -> bool:
    """Γραμμή-κεφαλίδα δικογράφου: κεφαλαία (εκτός αγκυλών εξατομίκευσης)."""
    stripped = _BRACKETS.sub("", line).strip()
    letters = [ch for ch in stripped if ch.isalpha()]
    if not letters:
        return False
    return all(ch == ch.upper() for ch in letters) and len(line) <= 100


def _build_letterhead(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    header = section.header
    accent = header.paragraphs[0]
    accent.alignment = WD_ALIGN_PARAGRAPH.CENTER
    accent.paragraph_format.space_after = Pt(2)
    accent_run = accent.add_run("─" * 8)
    _set_font(accent_run, name=BODY_FONT, size=8, color=GOLD, bold=True)

    name_par = header.add_paragraph()
    name_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_par.paragraph_format.space_after = Pt(0)
    name_run = name_par.add_run("SKOTANIS & ASSOCIATES")
    _set_font(name_run, name=DISPLAY_FONT, size=17, color=NAVY, bold=True)

    subtitle = header.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    subtitle_run = subtitle.add_run("E L I T E   L E G A L   A D V O C A C Y")
    _set_font(subtitle_run, name=BODY_FONT, size=7.5, color=GOLD_DARK, bold=True)
    _add_border(subtitle, "bottom", color="C6A75E", size=8, space=6)

    footer = section.footer
    contact = footer.paragraphs[0]
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_border(contact, "top", color="C6A75E", size=6, space=6)
    contact_run = contact.add_run(
        "Μύκονος: Αεροδρόμιο Μυκόνου · +30 2289 100 111    |    "
        "Αθήνα: Βαλαωρίτου 17, Κολωνάκι    |    "
        "skotanislaw.gr · ΑΜ ΔΣΣ 228"
    )
    _set_font(contact_run, name=BODY_FONT, size=7.5, color=SILVER)


def _render_body(document: Document, body: str) -> None:
    lines = body.splitlines()
    # Από τη γραμμή τόπου/ημερομηνίας και κάτω: μπλοκ υπογραφής, δεξιά στοίχιση.
    signature_start = next(
        (i for i, ln in enumerate(lines) if ln.strip().startswith("Μύκονος,")),
        len(lines),
    )

    blocks: list[tuple[list[str], bool]] = []
    current: list[str] = []
    in_signature = False
    for i, line in enumerate(lines):
        if i == signature_start:
            if current:
                blocks.append((current, in_signature))
                current = []
            in_signature = True
        if line.strip():
            current.append(line.strip())
        elif current:
            blocks.append((current, in_signature))
            current = []
    if current:
        blocks.append((current, in_signature))

    for block_lines, signature in blocks:
        for line in block_lines:
            par = document.add_paragraph()
            fmt = par.paragraph_format
            fmt.space_after = Pt(6)
            fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            fmt.line_spacing = 1.35
            if signature:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                run = par.add_run(line)
                _set_font(
                    run, name=BODY_FONT, size=11, color=NAVY,
                    bold=_is_display_line(line),
                )
            elif _is_display_line(line):
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                fmt.space_before = Pt(8)
                run = par.add_run(line)
                _set_font(run, name=BODY_FONT, size=11.5, color=NAVY, bold=True)
            else:
                par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                fmt.first_line_indent = Cm(0.75)
                run = par.add_run(line)
                _set_font(run, name=BODY_FONT, size=11, color=NAVY)


def build_exemplar_docx(exemplar: Exemplar) -> bytes:
    document = Document()
    _build_letterhead(document)

    notice = document.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.paragraph_format.space_after = Pt(14)
    notice_run = notice.add_run(
        f"ΥΠΟΔΕΙΓΜΑ ΓΡΑΦΕΙΟΥ — {exemplar.title} — ΠΡΟΣ ΕΞΑΤΟΜΙΚΕΥΣΗ"
    )
    _set_font(notice_run, name=BODY_FONT, size=8, color=GOLD_DARK, bold=True)

    _render_body(document, exemplar.body)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
