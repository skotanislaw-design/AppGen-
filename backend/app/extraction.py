"""Εξαγωγή κειμένου από PDF, DOCX και αρχεία απλού κειμένου."""

import io
from pathlib import PurePosixPath

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class ExtractionError(ValueError):
    pass


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("Το αρχείο PDF δεν μπόρεσε να αναγνωσθεί.") from exc
    if reader.is_encrypted:
        raise ExtractionError("Το PDF είναι κρυπτογραφημένο — αφαιρέστε τον κωδικό.")
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    text = "\n\n".join(p for p in pages if p)
    if not text.strip():
        raise ExtractionError(
            "Δεν εξήχθη κείμενο από το PDF — πιθανώς είναι σαρωμένο (εικόνα). "
            "Απαιτείται OCR ή επικόλληση του κειμένου."
        )
    return text


def _extract_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("Το αρχείο DOCX δεν μπόρεσε να αναγνωσθεί.") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    if not text.strip():
        raise ExtractionError("Το DOCX δεν περιέχει εξαγώγιμο κείμενο.")
    return text


def _extract_plain(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1253", "iso-8859-7"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Άγνωστη κωδικοποίηση χαρακτήρων του αρχείου.")


def extract_text(filename: str, data: bytes) -> str:
    """Επιστρέφει το κείμενο του αρχείου ή εγείρει ExtractionError."""
    ext = PurePosixPath(filename.lower()).suffix
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(
            "Μη υποστηριζόμενος τύπος αρχείου. Δεκτά: PDF, DOCX, TXT, MD."
        )
    if not data:
        raise ExtractionError("Το αρχείο είναι κενό.")
    if ext == ".pdf":
        return _extract_pdf(data)
    if ext == ".docx":
        return _extract_docx(data)
    return _extract_plain(data)
