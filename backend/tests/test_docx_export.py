import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from app.docx_export import build_exemplar_docx
from app.exemplars import EXEMPLARS


def _doc(doc_type: str) -> Document:
    data = build_exemplar_docx(EXEMPLARS[doc_type])
    assert data[:2] == b"PK"  # έγκυρο ZIP/OOXML
    return Document(io.BytesIO(data))


def test_every_exemplar_exports():
    for key in EXEMPLARS:
        _doc(key)


def test_letterhead_header_and_footer():
    doc = _doc("agogi")
    header_text = " ".join(p.text for p in doc.sections[0].header.paragraphs)
    assert "SKOTANIS & ASSOCIATES" in header_text
    footer_text = " ".join(p.text for p in doc.sections[0].footer.paragraphs)
    assert "Βαλαωρίτου 17" in footer_text
    assert "ΑΜ ΔΣΣ 228" in footer_text
    assert "2289 100 111" in footer_text


def test_gold_rules_present():
    doc = _doc("agogi")
    header_xml = doc.sections[0].header._element.xml
    footer_xml = doc.sections[0].footer._element.xml
    assert "C6A75E" in header_xml  # χρυσός κανόνας κεφαλίδας
    assert "C6A75E" in footer_xml  # χρυσή γραμμή υποσέλιδου


def test_body_contains_full_pleading():
    doc = _doc("exodiko")
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "ΕΞΩΔΙΚΗ ΔΗΛΩΣΗ" in text
    assert "ΧΡΗΣΤΟΣ ΣΚΟΤΑΝΗΣ" in text
    assert "[ΣΥΜΠΛΗΡΩΣΤΕ" in text


def test_customization_notice_present():
    doc = _doc("anakopi")
    first = doc.paragraphs[0].text
    assert "ΥΠΟΔΕΙΓΜΑ ΓΡΑΦΕΙΟΥ" in first
    assert "ΠΡΟΣ ΕΞΑΤΟΜΙΚΕΥΣΗ" in first


def test_signature_block_right_aligned():
    doc = _doc("agogi")
    sig = [p for p in doc.paragraphs if "ΧΡΗΣΤΟΣ ΣΚΟΤΑΝΗΣ" in p.text]
    assert sig and sig[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_display_lines_centered_and_body_justified():
    doc = _doc("agogi")
    by_text = {p.text: p for p in doc.paragraphs}
    assert by_text["ΑΓΩΓΗ"].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert by_text["ΚΑΤΑ"].alignment == WD_ALIGN_PARAGRAPH.CENTER
    body_pars = [
        p
        for p in doc.paragraphs
        if p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY and len(p.text) > 120
    ]
    assert body_pars, "Δεν βρέθηκαν πλήρως στοιχισμένες παράγραφοι σώματος"


def test_a4_page_setup():
    doc = _doc("agogi")
    section = doc.sections[0]
    assert abs(section.page_width.cm - 21.0) < 0.05
    assert abs(section.page_height.cm - 29.7) < 0.05


def test_greek_font_declared_on_runs():
    doc = _doc("agogi")
    run = next(
        r for p in doc.paragraphs for r in p.runs if "ΧΡΗΣΤΟΣ ΣΚΟΤΑΝΗΣ" in r.text
    )
    rfonts = run._element.get_or_add_rPr().find(qn("w:rFonts"))
    assert rfonts is not None
    assert rfonts.get(qn("w:hAnsi")) == "Georgia"


def test_endpoint_serves_docx():
    from fastapi.testclient import TestClient

    from app.main import app

    client = TestClient(app)
    res = client.get("/api/exemplars/agogi/docx")
    assert res.status_code == 200
    assert res.content[:2] == b"PK"
    assert "wordprocessingml" in res.headers["content-type"]
    assert "attachment" in res.headers["content-disposition"]
    assert client.get("/api/exemplars/unknown/docx").status_code == 404
